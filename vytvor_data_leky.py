"""Vytvoří fiktivní data výdejů léků pro další rok podle loňského souboru.

Vezme ukazkova_data_leky_pacienti.xlsx (rok 2025) a vytvoří stejný sešit pro
nový rok: stejné listy, formáty a vzorce, jen s novými smyšlenými výdeji.

- stávající pacienti pokračují se svými chronickými léky (podobně často jako loni),
- akutní a sezónní léky kopírují loňský průběh po měsících (víc v zimě),
- během roku přibude několik nových pacientů,
- ceny léků (modré vstupní buňky) se mírně zvýší.

Spuštění:
    python vytvor_data_leky.py                       # 2025 → 2026
    python vytvor_data_leky.py --zdroj data/ukazkova_data_leky_pacienti_2026.xlsx --rok 2027
"""

import argparse
import re
from copy import copy
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

def _kopiruj_styl(zdroj, cil) -> None:
    cil.font, cil.fill, cil.border = copy(zdroj.font), copy(zdroj.fill), copy(zdroj.border)
    cil.alignment, cil.number_format = copy(zdroj.alignment), zdroj.number_format


def nacti(wb) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    def list_na_df(nazev: str, sloupcu: int) -> pd.DataFrame:
        ws = wb[nazev]
        hlavicka = [c.value for c in ws[1]][:sloupcu]
        return pd.DataFrame([[c.value for c in r[:sloupcu]] for r in ws.iter_rows(min_row=2)
                             if r[0].value is not None], columns=hlavicka)

    leky = list_na_df("Léky", 9)
    pacienti = list_na_df("Pacienti", 7)
    vydeje = list_na_df("Výdeje", 6)[["Měsíc", "ID pacienta", "ID léku", "Počet balení"]]
    vydeje["Měsíc"] = pd.to_datetime(vydeje["Měsíc"])
    return leky, pacienti, vydeje


def vytvor_vydeje(leky: pd.DataFrame, pacienti: pd.DataFrame, vydeje: pd.DataFrame,
                  novi: pd.DataFrame, rok: int, rng: np.random.Generator) -> pd.DataFrame:
    typ = dict(zip(leky["ID léku"], leky["Typ užívání"]))
    chronicke_id = [i for i, t in typ.items() if t == "chronický"]
    podil_dvou = (vydeje["Počet balení"] > 1).mean()
    mesicu_loni = vydeje["Měsíc"].dt.month.nunique()
    radky = []

    # 1) chronické léky stávajících pacientů – stejná pravděpodobnost výdeje jako loni
    chron = vydeje[vydeje["ID léku"].map(typ) == "chronický"]
    prvni = chron.groupby("ID pacienta")["Měsíc"].min().dt.month
    for (pid, lid), n in chron.groupby(["ID pacienta", "ID léku"]).size().items():
        p = min(0.97, n / (mesicu_loni - prvni[pid] + 1))
        for m in range(1, 13):
            if rng.random() < p:
                radky.append((m, pid, lid))

    # 2) noví pacienti – 1 až 3 chronické léky od měsíce registrace
    for _, pac in novi.iterrows():
        for lid in rng.choice(chronicke_id, size=rng.integers(1, 4), replace=False):
            for m in range(pac["Datum registrace"].month, 13):
                if rng.random() < 0.85:
                    radky.append((m, pac["ID pacienta"], lid))

    # 3) akutní a sezónní léky – počty po měsících podle loňska, náhodní aktivní pacienti
    akutni = vydeje[vydeje["ID léku"].map(typ) != "chronický"]
    loni = akutni.groupby([akutni["Měsíc"].dt.month, "ID léku"]).size()
    vsichni = pd.concat([pacienti, novi])
    for m in range(1, 13):
        aktivni = vsichni.loc[pd.to_datetime(vsichni["Datum registrace"]) < pd.Timestamp(rok, m, 1)
                              + pd.offsets.MonthEnd(0), "ID pacienta"].tolist()
        for lid in (i for i, t in typ.items() if t != "chronický"):
            pocet = rng.poisson(loni.get((m, lid), 0) * 1.05)
            for pid in rng.choice(aktivni, size=min(pocet, len(aktivni)), replace=False):
                radky.append((m, pid, lid))

    df = pd.DataFrame(radky, columns=["mesic", "ID pacienta", "ID léku"]).drop_duplicates()
    df["Počet balení"] = np.where(rng.random(len(df)) < podil_dvou, 2, 1)
    df["Měsíc"] = pd.to_datetime({"year": rok, "month": df["mesic"], "day": 1})
    return df.sort_values(["Měsíc", "ID pacienta", "ID léku"])[["Měsíc", "ID pacienta", "ID léku", "Počet balení"]]


def novi_pacienti(pacienti: pd.DataFrame, pocet: int, rok: int, rng: np.random.Generator) -> pd.DataFrame:
    posledni = int(pacienti["ID pacienta"].str[1:].astype(int).max())
    registrace = sorted(pd.Timestamp(rok, 1, 1) + pd.to_timedelta(rng.integers(0, 330, pocet), unit="D"))
    return pd.DataFrame({
        "ID pacienta": [f"P{posledni + i + 1:04d}" for i in range(pocet)],
        "Pohlaví": rng.choice(["M", "Ž"], pocet),
        "Rok narození": rng.integers(1945, 2005, pocet),
        "Kraj": rng.choice(pacienti["Kraj"], pocet),
        "Ošetřující lékař": rng.choice(pacienti["Ošetřující lékař"], pocet),
        "Datum registrace": registrace,
    })


def uloz(wb, novi: pd.DataFrame, vydeje: pd.DataFrame, rok_zdroj: int, rok: int,
         zdrazeni: float, cesta: Path) -> None:
    # Info – texty s rokem a počtem pacientů
    ws = wb["Info"]
    pocet_pac = wb["Pacienti"].max_row - 1 + len(novi)
    for r in ws.iter_rows():
        for c in r:
            if isinstance(c.value, str):
                c.value = c.value.replace(str(rok_zdroj), str(rok))
                c.value = re.sub(r"\d+ pacientů", f"{pocet_pac} pacientů", c.value)

    # Léky – mírné zdražení (modré vstupní buňky)
    ws = wb["Léky"]
    for r in range(2, ws.max_row + 1):
        c = ws.cell(r, 8)
        if isinstance(c.value, (int, float)):
            c.value = int(round(c.value * (1 + zdrazeni)))

    # Pacienti – věk k novému roku + noví pacienti
    ws = wb["Pacienti"]
    ws["D1"] = f"Věk ({rok})"
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 4).value = f"={rok}-C{r}"
    vzor = [ws.cell(2, s) for s in range(1, 8)]
    for _, p in novi.iterrows():
        r = ws.max_row + 1
        hodnoty = [p["ID pacienta"], p["Pohlaví"], int(p["Rok narození"]), f"={rok}-C{r}", p["Kraj"],
                   p["Ošetřující lékař"], p["Datum registrace"].to_pydatetime()]
        for s, h in enumerate(hodnoty, start=1):
            ws.cell(r, s, h)
            _kopiruj_styl(vzor[s - 1], ws.cell(r, s))

    # Výdeje – přepsat řádky, název/skupina/cena zůstávají vzorcem podle ID léku
    ws = wb["Výdeje"]
    vzor = [ws.cell(2, s) for s in range(1, 9)]
    styly = [(c.font, c.fill, c.border, c.alignment, c.number_format) for c in vzor]  # před smazáním řádků
    ws.delete_rows(2, ws.max_row)
    n_leku = wb["Léky"].max_row
    for i, v in enumerate(vydeje.itertuples(index=False), start=2):
        hledej = f"MATCH(C{i},Léky!$A$2:$A${n_leku},0))"
        hodnoty = [v[0].to_pydatetime(), v[1], v[2],
                   f"=INDEX(Léky!$B$2:$B${n_leku},{hledej}", f"=INDEX(Léky!$E$2:$E${n_leku},{hledej}",
                   int(v[3]), f"=INDEX(Léky!$H$2:$H${n_leku},{hledej}", f"=F{i}*G{i}"]
        for s, h in enumerate(hodnoty, start=1):
            c = ws.cell(i, s, h)
            f, fi, b, a, nf = styly[s - 1]
            c.font, c.fill, c.border, c.alignment, c.number_format = copy(f), copy(fi), copy(b), copy(a), nf
    posledni = len(vydeje) + 1
    ws.auto_filter.ref = f"A1:H{posledni}"

    # Přehled po měsících – rozsahy SUMIFS a rok v DATE()
    ws = wb["Přehled po měsících"]
    for r in ws.iter_rows():
        for c in r:
            if isinstance(c.value, str) and c.value.startswith("=SUMIFS"):
                c.value = re.sub(r"(Výdeje!\$[A-H]\$2:\$[A-H]\$)\d+", rf"\g<1>{posledni}", c.value)
                c.value = c.value.replace(f"DATE({rok_zdroj},", f"DATE({rok},")

    cesta.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cesta)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--zdroj", default="data/ukazkova_data_leky_pacienti.xlsx")
    parser.add_argument("--rok", type=int, default=2026)
    parser.add_argument("--novych-pacientu", type=int, default=6)
    parser.add_argument("--zdrazeni", type=float, default=0.03, help="růst cen, 0.03 = +3 %%")
    parser.add_argument("--seed", type=int, default=2026)
    parser.add_argument("-o", "--vystup", help="výchozí data/ukazkova_data_leky_pacienti_ROK.xlsx")
    args = parser.parse_args()

    rng = np.random.default_rng(args.seed)
    wb = load_workbook(args.zdroj)
    leky, pacienti, vydeje = nacti(wb)
    rok_zdroj = int(vydeje["Měsíc"].dt.year.mode()[0])

    novi = novi_pacienti(pacienti, args.novych_pacientu, args.rok, rng)
    nove_vydeje = vytvor_vydeje(leky, pacienti, vydeje, novi, args.rok, rng)
    cesta = Path(args.vystup or f"data/ukazkova_data_leky_pacienti_{args.rok}.xlsx")
    uloz(wb, novi, nove_vydeje, rok_zdroj, args.rok, args.zdrazeni, cesta)
    print(f"Uloženo: {cesta} – {len(nove_vydeje)} výdejů (loni {len(vydeje)}), "
          f"{len(pacienti) + len(novi)} pacientů (+{len(novi)} nových)")


if __name__ == "__main__":
    main()
