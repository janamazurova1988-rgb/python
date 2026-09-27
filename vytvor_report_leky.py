"""Vytvoří měsíční report výdejů léků (stejná struktura jako report_2025_12.xlsx).

Načte sešit s výdeji (listy Léky, Pacienti, Výdeje), spočítá srovnání vybraného
měsíce s předchozím a uloží Excel s listy Report, Grafy a Výdeje v měsíci.

Spuštění:
    python vytvor_report_leky.py                     # prosinec 2026 z data/ukazkova_data_leky_pacienti_2026.xlsx
    python vytvor_report_leky.py --rok 2025 --mesic 12
    python vytvor_report_leky.py --data moje_vydeje.xlsx --rok 2026 --mesic 6
"""

import argparse
import tempfile
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import matplotlib.ticker as mticker  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.drawing.image import Image as XlImage  # noqa: E402
from openpyxl.formatting.rule import CellIsRule  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402

MESICE = ["leden", "únor", "březen", "duben", "květen", "červen",
          "červenec", "srpen", "září", "říjen", "listopad", "prosinec"]
MESICE_KRATCE = ["led", "úno", "bře", "dub", "kvě", "čvn", "čvc", "srp", "zář", "říj", "lis", "pro"]

MODRA_TMAVA = "1F4E78"
FMT_CISLO = "#,##0;\\-#,##0;\\-"
FMT_PROC = "\\+0.0%;\\-0.0%;0.0%"


# ------------------------------------------------------------------- data
def nacti_vydeje(cesta: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Načte sešit s listy Léky, Pacienti a Výdeje; název, cenu a kraj doplní podle ID."""
    listy = pd.read_excel(cesta, sheet_name=["Léky", "Pacienti", "Výdeje"])
    leky, pacienti = listy["Léky"], listy["Pacienti"]
    v = listy["Výdeje"][["Měsíc", "ID pacienta", "ID léku", "Počet balení"]].dropna()
    v["Měsíc"] = pd.to_datetime(v["Měsíc"])
    v = v.merge(pacienti[["ID pacienta", "Kraj"]], on="ID pacienta", how="left")
    v = v.merge(leky[["ID léku", "Název", "Cena za balení (Kč)"]], on="ID léku", how="left")
    v["Tržby (Kč)"] = v["Počet balení"] * v["Cena za balení (Kč)"]
    sloupce = ["Měsíc", "ID pacienta", "Kraj", "ID léku", "Název", "Počet balení",
               "Cena za balení (Kč)", "Tržby (Kč)"]
    return v[sloupce].sort_values(["Měsíc", "ID pacienta", "ID léku"]).reset_index(drop=True), pacienti


def ukazatele(df: pd.DataFrame, pacienti: pd.DataFrame, datum: pd.Timestamp) -> list[int]:
    m = df[df["Měsíc"] == datum]
    registrace = pd.to_datetime(pacienti["Datum registrace"]).dt.to_period("M")
    novi = int((registrace == datum.to_period("M")).sum())
    return [int(m["Tržby (Kč)"].sum()), int(m["Počet balení"].sum()), len(m), m["ID pacienta"].nunique(), novi]


def srovnani(df: pd.DataFrame, podle: str, akt: pd.Timestamp, pred: pd.Timestamp) -> pd.DataFrame:
    t = df[df["Měsíc"].isin([akt, pred])].pivot_table(
        index=podle, columns="Měsíc", values=["Tržby (Kč)", "Počet balení"], aggfunc="sum", fill_value=0)
    out = pd.DataFrame({
        "trzby_akt": t[("Tržby (Kč)", akt)], "trzby_pred": t[("Tržby (Kč)", pred)],
        "baleni_akt": t[("Počet balení", akt)], "baleni_pred": t[("Počet balení", pred)],
    }).astype(int)
    return out.sort_values("trzby_akt", ascending=False)


# ------------------------------------------------------------------ grafy
def _osy(ax) -> None:
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.set_axisbelow(True)


def graf_mesicni_trzby(df: pd.DataFrame, mesic: int, cesta: Path) -> None:
    trzby = df.groupby(df["Měsíc"].dt.month)["Tržby (Kč)"].sum().reindex(range(1, 13), fill_value=0)
    barvy = ["#2a78d6" if m == mesic else "#c8c7c2" for m in trzby.index]
    fig, ax = plt.subplots(figsize=(9, 4), dpi=130)
    ax.bar(MESICE_KRATCE, trzby.values, color=barvy, width=0.65)
    ax.grid(axis="y", color="#e4e3df")
    ax.yaxis.set_major_formatter(mticker.FuncFormatter(lambda v, _: f"{v:,.0f}".replace(",", " ")))
    _osy(ax)
    ax.set_title("Tržby po měsících (Kč)", loc="left", fontweight="bold")
    fig.tight_layout()
    fig.savefig(cesta, facecolor="white")
    plt.close(fig)


def graf_zmena_leku(podle_leku: pd.DataFrame, nazev_akt: str, nazev_pred: str, cesta: Path) -> None:
    zmena = (podle_leku["trzby_akt"] / podle_leku["trzby_pred"].replace(0, np.nan) - 1).fillna(0) * 100
    zmena = zmena.sort_values()
    fig, ax = plt.subplots(figsize=(9, 4.5), dpi=130)
    barvy = ["#1baf7a" if v > 0 else "#e34948" for v in zmena]
    ax.barh(zmena.index, zmena.values, color=barvy, height=0.6)
    ax.axvline(0, color="#52514e", linewidth=1)
    rozpeti = max(zmena.max(), 10) - min(zmena.min(), -10)
    ax.set_xlim(min(zmena.min(), -10) - rozpeti * 0.12, max(zmena.max(), 10) + rozpeti * 0.12)
    for y, v in enumerate(zmena.values):
        ax.text(v + (1 if v >= 0 else -1) * rozpeti * 0.01, y, f"{v:+.0f}%", va="center",
                ha="left" if v >= 0 else "right", fontsize=9)
    ax.xaxis.set_major_formatter(mticker.FuncFormatter(lambda v, _: f"{v:+.0f} %"))
    ax.grid(axis="x", color="#e4e3df")
    _osy(ax)
    ax.set_title(f"Změna tržeb podle léku: {nazev_akt} vs. {nazev_pred}", loc="left", fontweight="bold")
    fig.tight_layout()
    fig.savefig(cesta, facecolor="white")
    plt.close(fig)


# ------------------------------------------------------------------ Excel
def uloz_report(df: pd.DataFrame, pacienti: pd.DataFrame, zdroj: str, rok: int, mesic: int, cesta: Path) -> None:
    akt = pd.Timestamp(rok, mesic, 1)
    pred = akt - pd.DateOffset(months=1)
    n_akt = f"{MESICE[akt.month - 1]} {akt.year}"
    n_pred = f"{MESICE[pred.month - 1]} {pred.year}"

    bila_tucne = Font(bold=True, color="FFFFFF")
    hlavicka = PatternFill("solid", fgColor=MODRA_TMAVA)
    cervena = PatternFill(bgColor="FBE3E3")
    zelena = PatternFill(bgColor="E1F5EC")

    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws["A1"] = f"Měsíční report výdejů – {n_akt}"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = f"Zdroj: {zdroj} · srovnání s měsícem {n_pred} · zvýrazněny změny nad ±20%"

    def zahlavi(radek: int, hodnoty: list[str]) -> None:
        for sl, h in enumerate(hodnoty, start=1):
            c = ws.cell(radek, sl, h)
            c.font, c.fill = bila_tucne, hlavicka

    def zvyrazni(rozsah: str) -> None:
        ws.conditional_formatting.add(rozsah, CellIsRule(operator="lessThan", formula=["-0.2"], fill=cervena))
        ws.conditional_formatting.add(rozsah, CellIsRule(operator="greaterThan", formula=["0.2"], fill=zelena))

    def radek_hodnot(r: int, popis: str, hodnoty: list, tucne: bool = False) -> None:
        ws.cell(r, 1, popis)
        for sl, h in enumerate(hodnoty, start=2):
            c = ws.cell(r, sl, h)
            c.number_format = FMT_PROC if sl == 5 else FMT_CISLO
        if tucne:
            for sl in range(1, len(hodnoty) + 2):
                ws.cell(r, sl).font = Font(bold=True)

    # Hlavní ukazatele
    ws["A4"] = "Hlavní ukazatele"
    ws["A4"].font = Font(bold=True, size=12)
    zahlavi(5, ["Ukazatel", n_akt, n_pred, "Rozdíl", "Změna %"])
    nazvy = ["Tržby (Kč)", "Vydaná balení", "Počet výdejů", "Aktivní pacienti", "Noví pacienti"]
    for i, (nazev, a, p) in enumerate(zip(nazvy, ukazatele(df, pacienti, akt), ukazatele(df, pacienti, pred))):
        r = 6 + i
        radek_hodnot(r, nazev, [a, p, f"=B{r}-C{r}", f'=IF(C{r}=0,"",B{r}/C{r}-1)'])
    zvyrazni("E6:E10")

    # Tabulky podle léku a kraje
    def tabulka(start: int, titulek: str, prvni: str, data: pd.DataFrame) -> int:
        ws.cell(start, 1, titulek).font = Font(bold=True, size=12)
        zahlavi(start + 1, [prvni, f"Tržby {n_akt}", f"Tržby {n_pred}", "Rozdíl (Kč)", "Změna %",
                            f"Balení {n_akt}", f"Balení {n_pred}"])
        r0 = start + 2
        for i, (nazev, v) in enumerate(data.iterrows()):
            r = r0 + i
            radek_hodnot(r, nazev, [v.trzby_akt, v.trzby_pred, f"=B{r}-C{r}", f'=IF(C{r}=0,"",B{r}/C{r}-1)',
                                    v.baleni_akt, v.baleni_pred])
        r1 = r0 + len(data) - 1
        rc = r1 + 1
        radek_hodnot(rc, "Celkem", [f"=SUM(B{r0}:B{r1})", f"=SUM(C{r0}:C{r1})", f"=SUM(D{r0}:D{r1})",
                                    f'=IF(C{rc}=0,"",B{rc}/C{rc}-1)', f"=SUM(F{r0}:F{r1})",
                                    f"=SUM(G{r0}:G{r1})"], tucne=True)
        zvyrazni(f"E{r0}:E{r1}")
        return rc

    podle_leku = srovnani(df, "Název", akt, pred)
    konec = tabulka(13, "Podle léku", "Lék", podle_leku)
    tabulka(konec + 3, "Podle kraje", "Kraj", srovnani(df, "Kraj", akt, pred))
    ws.column_dimensions["A"].width = 22
    for sl in "BCDEFG":
        ws.column_dimensions[sl].width = 20

    # Grafy
    wg = wb.create_sheet("Grafy")
    wg["A1"] = f"Grafy – {n_akt}"
    wg["A1"].font = Font(bold=True, size=14)
    with tempfile.TemporaryDirectory() as tmp:
        g1, g2 = Path(tmp) / "trzby.png", Path(tmp) / "zmena.png"
        graf_mesicni_trzby(df, mesic, g1)
        graf_zmena_leku(podle_leku, n_akt, n_pred, g2)
        for obr, bunka in ((g1, "A2"), (g2, "A28")):
            img = XlImage(str(obr))
            img.width, img.height = 1170, img.height * 1170 / img.width
            wg.add_image(img, bunka)

        # Výdeje v měsíci
        wv = wb.create_sheet("Výdeje v měsíci")
        vydeje = df[df["Měsíc"] == akt]
        wv.append(list(vydeje.columns))
        for c in wv[1]:
            c.font, c.fill = bila_tucne, hlavicka
        for radek in vydeje.itertuples(index=False):
            wv.append([radek[0].to_pydatetime(), *radek[1:]])
        for (c,) in wv.iter_rows(min_row=2, max_col=1):
            c.number_format = "mm/yyyy"
        wv.freeze_panes = "A2"
        for sl, sirka in zip("ABCDEFGH", [10, 12, 17, 9, 22, 12, 14, 12]):
            wv.column_dimensions[sl].width = sirka

        cesta.parent.mkdir(parents=True, exist_ok=True)
        wb.save(cesta)  # obrázky se načítají při ukládání, proto ještě uvnitř bloku


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--data", help="sešit s výdeji (výchozí data/ukazkova_data_leky_pacienti_ROK.xlsx)")
    parser.add_argument("--rok", type=int, default=2026)
    parser.add_argument("--mesic", type=int, default=12, choices=range(2, 13), metavar="2–12")
    parser.add_argument("-o", "--vystup", help="výstupní soubor (výchozí reporty/report_ROK_MM.xlsx)")
    args = parser.parse_args()

    cesta = Path(args.vystup or f"reporty/report_{args.rok}_{args.mesic:02d}.xlsx")
    data = Path(args.data or f"data/ukazkova_data_leky_pacienti_{args.rok}.xlsx")
    df, pacienti = nacti_vydeje(data)
    uloz_report(df, pacienti, data.name, args.rok, args.mesic, cesta)
    print(f"Uloženo: {cesta} ({len(df)} výdejů za rok, {df['ID pacienta'].nunique()} pacientů)")


if __name__ == "__main__":
    main()
