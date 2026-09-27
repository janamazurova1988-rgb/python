"""Měsíční analýza prodejů z Excelu.

Načte Excel se sloupci datum, produkt, region, hodnota, spočítá měsíční souhrny
podle produktu a regionu, vykreslí grafy a uloží:

    vystup/souhrn_prodeju.xlsx              – souhrnné tabulky + list s grafy
    vystup/prodeje_podle_produktu.png       – měsíční vývoj po produktech
    vystup/prodeje_podle_regionu.png        – měsíční vývoj po regionech
    vystup/produkt_region_heatmapa.png      – součet za období: produkt × region

Spuštění:
    python analyza_prodeju.py                                  # ukázková data
    python analyza_prodeju.py moje_data.xlsx -o vysledky       # vlastní soubor
    python analyza_prodeju.py data.xlsx --list Prodeje         # jiný list
"""

import argparse
from pathlib import Path

import matplotlib

matplotlib.use("Agg")  # kreslení do souboru, bez okna
import matplotlib.pyplot as plt  # noqa: E402
import matplotlib.ticker as mticker  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl.drawing.image import Image as XlImage  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

POVINNE_SLOUPCE = ["datum", "produkt", "region", "hodnota"]
NAZVY_V_EXCELU = {"mesic": "Měsíc", "produkt": "Produkt", "region": "Region",
                  "soucet": "Součet", "prumer": "Průměr", "pocet_zaznamu": "Počet záznamů"}

# Barvy sérií – pevné pořadí, ověřené pro barvoslepost (sousední dvojice).
BARVY = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300", "#4a3aa7", "#e34948"]
TEXT = "#0b0b0b"
TEXT_2 = "#52514e"
MRIZKA = "#e4e3df"

MESICE = ["led", "úno", "bře", "dub", "kvě", "čvn", "čvc", "srp", "zář", "říj", "lis", "pro"]


# ---------------------------------------------------------------- načtení dat
def nacti_data(cesta: Path, list_: str | int = 0) -> pd.DataFrame:
    df = pd.read_excel(cesta, sheet_name=list_)
    # názvy sloupců bez ohledu na velikost písmen a mezery ("Datum ", "PRODUKT")
    df.columns = [str(c).strip().lower() for c in df.columns]
    chybi = [c for c in POVINNE_SLOUPCE if c not in df.columns]
    if chybi:
        raise ValueError(f"V souboru chybí sloupce: {', '.join(chybi)} (nalezeno: {list(df.columns)})")

    df = df[POVINNE_SLOUPCE].copy()
    df["datum"] = pd.to_datetime(df["datum"], errors="coerce", dayfirst=True)
    df["hodnota"] = pd.to_numeric(df["hodnota"], errors="coerce")
    df["produkt"] = df["produkt"].astype(str).str.strip()
    df["region"] = df["region"].astype(str).str.strip()

    vadne = df["datum"].isna() | df["hodnota"].isna()
    if vadne.any():
        print(f"Upozornění: vynechávám {vadne.sum()} řádků s chybným datem nebo hodnotou.")
        df = df[~vadne]
    if df.empty:
        raise ValueError("Po vyčištění nezůstala žádná data.")

    # funguje pro denní i měsíční záznamy – vše se sloučí do kalendářních měsíců
    df["mesic"] = df["datum"].dt.to_period("M")
    return df


# ------------------------------------------------------------------ souhrny
def spocitej_souhrny(df: pd.DataFrame) -> dict[str, pd.DataFrame]:
    detail = (
        df.groupby(["mesic", "produkt", "region"])["hodnota"]
        .agg(soucet="sum", prumer="mean", pocet_zaznamu="count")
        .reset_index()
    )
    detail["prumer"] = detail["prumer"].round(1)

    def kontingence(sloupce: str) -> pd.DataFrame:
        t = df.pivot_table(index="mesic", columns=sloupce, values="hodnota", aggfunc="sum", fill_value=0)
        t = t[t.sum().sort_values(ascending=False).index]  # největší vlevo
        t["Celkem"] = t.sum(axis=1)
        return t

    mesic_produkt = kontingence("produkt")
    mesic_region = kontingence("region")

    produkt_region = df.pivot_table(index="produkt", columns="region", values="hodnota", aggfunc="sum", fill_value=0)
    produkt_region = produkt_region.loc[
        produkt_region.sum(axis=1).sort_values(ascending=False).index,
        produkt_region.sum().sort_values(ascending=False).index,
    ]
    produkt_region["Celkem"] = produkt_region.sum(axis=1)
    produkt_region.loc["Celkem"] = produkt_region.sum()

    return {
        "Detail": detail,
        "Mesic x Produkt": mesic_produkt,
        "Mesic x Region": mesic_region,
        "Produkt x Region": produkt_region,
    }


# -------------------------------------------------------------------- grafy
def _popisek_mesice(p: pd.Period) -> str:
    return f"{MESICE[p.month - 1]}\n{p.year}" if p.month == 1 else MESICE[p.month - 1]


def _styl_os(ax) -> None:
    for strana in ("top", "right", "left"):
        ax.spines[strana].set_visible(False)
    ax.spines["bottom"].set_color(TEXT_2)
    ax.tick_params(colors=TEXT_2, length=0)
    ax.grid(axis="y", color=MRIZKA, linewidth=0.8)
    ax.set_axisbelow(True)
    ax.yaxis.set_major_formatter(mticker.FuncFormatter(lambda v, _: f"{v:,.0f}".replace(",", " ")))


def graf_vyvoje(tabulka: pd.DataFrame, titulek: str, cesta: Path) -> None:
    """Spojnicový graf měsíčního vývoje, jedna čára na sérii."""
    serie = [c for c in tabulka.columns if c != "Celkem"]
    if len(serie) > len(BARVY):  # víc sérií než barev → nejmenší sloučit do „Ostatní“
        hlavni = serie[: len(BARVY) - 1]
        tabulka = tabulka[hlavni].assign(Ostatní=tabulka[serie[len(BARVY) - 1:]].sum(axis=1))
        serie = list(tabulka.columns)

    x = range(len(tabulka))
    fig, ax = plt.subplots(figsize=(11, 5.5), dpi=150)
    for i, s in enumerate(serie):
        ax.plot(x, tabulka[s], color=BARVY[i], linewidth=2, marker="o", markersize=4, label=s)

    # přímé popisky na konci čar, jen pokud jich je málo a nepřekrývají se
    if len(serie) <= 5:
        posledni = tabulka[serie].iloc[-1].sort_values()
        rozsah = ax.get_ylim()[1] - ax.get_ylim()[0]
        y_prev = -float("inf")
        for s, y in posledni.items():
            y = max(y, y_prev + rozsah * 0.045)
            ax.annotate(s, (len(tabulka) - 1, y), xytext=(8, 0), textcoords="offset points",
                        va="center", fontsize=9, color=TEXT)
            y_prev = y
        ax.set_xlim(-0.3, len(tabulka) - 1 + 2.2)

    ax.set_xticks(list(x), [_popisek_mesice(p) for p in tabulka.index])
    ax.set_ylim(bottom=0)
    _styl_os(ax)
    ax.set_title(titulek, loc="left", fontsize=14, color=TEXT, pad=28)
    ax.legend(loc="lower left", bbox_to_anchor=(0, 1.0), ncol=len(serie), frameon=False,
              fontsize=9, labelcolor=TEXT, handlelength=1.5)
    fig.tight_layout()
    fig.savefig(cesta, facecolor="white")
    plt.close(fig)


def graf_heatmapa(produkt_region: pd.DataFrame, obdobi: str, cesta: Path) -> None:
    """Heatmapa součtů produkt × region za celé období (bez řádku/sloupce Celkem)."""
    data = produkt_region.drop(index="Celkem", columns="Celkem")
    fig, ax = plt.subplots(figsize=(8, 0.6 * len(data) + 2), dpi=150)
    im = ax.imshow(data.values, cmap="Blues", aspect="auto", vmin=0)

    prah = data.values.max() * 0.6  # na tmavých buňkách bílý text
    for (r, c), v in pd.DataFrame(data.values).stack().items():
        ax.text(c, r, f"{v:,.0f}".replace(",", " "), ha="center", va="center", fontsize=9,
                color="white" if v > prah else TEXT)

    ax.set_xticks(range(data.shape[1]), data.columns)
    ax.set_yticks(range(data.shape[0]), data.index)
    ax.tick_params(length=0, colors=TEXT)
    ax.xaxis.tick_top()
    for s in ax.spines.values():
        s.set_visible(False)
    ax.set_title(f"Prodeje za celé období ({obdobi}): produkt × region", loc="left", fontsize=13,
                 color=TEXT, pad=30)
    cb = fig.colorbar(im, ax=ax, shrink=0.8)
    cb.outline.set_visible(False)
    cb.ax.tick_params(colors=TEXT_2, length=0)
    cb.ax.yaxis.set_major_formatter(mticker.FuncFormatter(lambda v, _: f"{v:,.0f}".replace(",", " ")))
    fig.tight_layout()
    fig.savefig(cesta, facecolor="white")
    plt.close(fig)


# -------------------------------------------------------------- uložení Excelu
def uloz_excel(souhrny: dict[str, pd.DataFrame], obrazky: list[Path], cesta: Path) -> None:
    hlavicka_font = Font(bold=True, color="FFFFFF")
    hlavicka_vypln = PatternFill("solid", fgColor="2A78D6")
    celkem_font = Font(bold=True)

    with pd.ExcelWriter(cesta, engine="openpyxl") as writer:
        for nazev, tabulka in souhrny.items():
            t = tabulka.copy()
            if "mesic" in t.columns:
                t["mesic"] = t["mesic"].astype(str)
            if isinstance(t.index, pd.PeriodIndex):
                t.index = t.index.astype(str)
            t.columns.name = None
            t.index.name = NAZVY_V_EXCELU.get(t.index.name, t.index.name)
            t = t.rename(columns=NAZVY_V_EXCELU)
            t.to_excel(writer, sheet_name=nazev, index=nazev != "Detail")

            ws = writer.sheets[nazev]
            ws.freeze_panes = "B2"
            for bunka in ws[1]:
                bunka.font = hlavicka_font
                bunka.fill = hlavicka_vypln
            for sloupec in ws.iter_cols(min_row=2):
                for bunka in sloupec:
                    if isinstance(bunka.value, (int, float)):
                        bunka.number_format = "#,##0" if float(bunka.value).is_integer() else "#,##0.0"
            # zvýraznit řádek a sloupec Celkem
            for radek in ws.iter_rows():
                if radek[0].value == "Celkem":
                    for bunka in radek:
                        bunka.font = celkem_font
            for bunka in ws[1]:
                if bunka.value == "Celkem":
                    for b in ws[bunka.column_letter][1:]:
                        b.font = celkem_font
            # šířka sloupců podle obsahu
            for i, sloupec in enumerate(ws.iter_cols(), start=1):
                sirka = max(len(str(b.value)) if b.value is not None else 0 for b in sloupec)
                ws.column_dimensions[get_column_letter(i)].width = min(max(sirka + 2, 10), 40)

        ws = writer.book.create_sheet("Grafy")
        radek = 1
        for obr in obrazky:
            img = XlImage(str(obr))
            img.width, img.height = img.width * 0.45, img.height * 0.45
            ws.add_image(img, f"A{radek}")
            radek += int(img.height / 20) + 2


# --------------------------------------------------------------------- main
def main() -> None:
    parser = argparse.ArgumentParser(description="Měsíční souhrny prodejů z Excelu + grafy.")
    parser.add_argument("vstup", nargs="?", default="data/ukazkova_data_prodeje.xlsx",
                        help="vstupní Excel (výchozí: ukázková data)")
    parser.add_argument("-o", "--vystup", default="vystup", help="složka pro výsledky")
    parser.add_argument("--list", default=0, help="název nebo pořadí listu (výchozí první)")
    args = parser.parse_args()

    list_ = int(args.list) if str(args.list).isdigit() else args.list
    vstup = Path(args.vstup)
    vystup = Path(args.vystup)
    vystup.mkdir(parents=True, exist_ok=True)

    df = nacti_data(vstup, list_)
    od, do = df["mesic"].min(), df["mesic"].max()
    print(f"Načteno {len(df)} řádků, období {od}–{do}, "
          f"{df['produkt'].nunique()} produktů, {df['region'].nunique()} regionů.")

    souhrny = spocitej_souhrny(df)

    obrazky = [
        vystup / "prodeje_podle_produktu.png",
        vystup / "prodeje_podle_regionu.png",
        vystup / "produkt_region_heatmapa.png",
    ]
    graf_vyvoje(souhrny["Mesic x Produkt"], "Měsíční prodeje podle produktu", obrazky[0])
    graf_vyvoje(souhrny["Mesic x Region"], "Měsíční prodeje podle regionu", obrazky[1])
    graf_heatmapa(souhrny["Produkt x Region"], f"{od}–{do}", obrazky[2])

    excel = vystup / "souhrn_prodeju.xlsx"
    uloz_excel(souhrny, obrazky, excel)

    print(f"\nCelkem za období: {df['hodnota'].sum():,.0f}".replace(",", " "))
    print("\nTop produkty:")
    print(souhrny["Produkt x Region"]["Celkem"].drop("Celkem").to_string())
    print(f"\nUloženo:\n  {excel}")
    for o in obrazky:
        print(f"  {o}")


if __name__ == "__main__":
    main()
